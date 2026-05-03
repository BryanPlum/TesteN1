from openpyxl import Workbook, load_workbook
import os


DB_FILE = "data/db.xlsx"


def init_db():
    if not os.path.exists("data"):
        os.makedirs("data")

    if not os.path.exists(DB_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Users"
        ws.append(["Usuario", "Senha", "Case1_1", "Case1_2"])
        wb.save(DB_FILE)


def save_user(email, password):
    wb = load_workbook(DB_FILE)
    ws = wb.active
    ws.append([email, password])
    wb.save(DB_FILE)

init_db()

def get_user(target_email):
    wb = load_workbook(DB_FILE)
    ws = wb.active

    for row in ws.iter_rows(min_row=2, values_only=True):
        email = row[0]
        senha = row[1]

        if email == target_email:
            return senha

    return None

#________ATUALIZAR SENHA DE USUÁRIO_______
def update_password(email, new_password):
    wb = load_workbook(DB_FILE)
    ws = wb.active

    for row in ws.iter_rows(min_row=2):
        if row[0].value == email:
            row[1].value = new_password
            break

    wb.save(DB_FILE)

#______FUNÇÃO SALVAR RESPOSTA NO BANCO______
def save_case1_answers(email, answer1, answer2):
    wb = load_workbook(DB_FILE)
    ws = wb.active

    for row in ws.iter_rows(min_row=2):
        if row[0].value == email:
            row[2].value = answer1  # Case1_1
            row[3].value = answer2  # Case1_2
            break

    wb.save(DB_FILE)    